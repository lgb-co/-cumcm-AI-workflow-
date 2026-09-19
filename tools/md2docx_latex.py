"""Markdown → DOCX（数学公式一律保留为 LaTeX 文本，不插入任何公式图片）。

用途：给作者"能看、能改公式"的查看版导出。正式提交版排版不适用本工具。
用法：python md2docx_latex.py <输入.md> <输出.docx>
"""
from __future__ import annotations

import io
import re
import sys

try:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt
except ModuleNotFoundError:
    sys.stderr.write(
        "缺少 python-docx，无法导出 DOCX。\n"
        "请先运行本工具同目录的 setup-md2docx.ps1 建立导出环境。\n"
        "当前解释器：" + sys.executable + "\n"
    )
    raise SystemExit(3)


def set_font(run, size=12, bold=False, east="宋体", ascii_font="Times New Roman"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = ascii_font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east)


def add_para(doc, text, size=12, bold=False, align=None, east="宋体", mono=False, indent=False, quote=False):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_after = Pt(6)
    pf.line_spacing = 1.25
    if indent:
        pf.first_line_indent = Cm(0.85)
    run = p.add_run(text)
    if mono:
        set_font(run, 10, bold, east="宋体", ascii_font="Consolas")
    else:
        set_font(run, size, bold, east=east)
    if quote:
        run.font.italic = True
    return p


def split_row(line: str) -> list[str]:
    """按 | 分列，但保护 $...$ 数学里的竖线（如 \\right|_{r=0}）。"""
    s = line.strip()
    if s.startswith("|"):
        s = s[1:]
    if s.endswith("|"):
        s = s[:-1]
    cells, cur, in_math = [], [], False
    for ch in s:
        if ch == "$":
            in_math = not in_math
            cur.append(ch)
        elif ch == "|" and not in_math:
            cells.append("".join(cur).strip())
            cur = []
        else:
            cur.append(ch)
    cells.append("".join(cur).strip())
    return cells


def add_table(doc, rows):
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    table = doc.add_table(rows=len(rows), cols=width)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j, cell in enumerate(row):
            c = table.cell(i, j)
            c.text = ""
            para = c.paragraphs[0]
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
            run = para.add_run(cell)
            set_font(run, 10.5, bold=(i == 0))


def main() -> int:
    src_path, out_path = sys.argv[1], sys.argv[2]
    lines = io.open(src_path, encoding="utf-8").read().splitlines()
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.4)
    sec.top_margin = sec.bottom_margin = Cm(2.4)

    i, table_buf, in_code = 0, [], False
    while i < len(lines):
        raw = lines[i].rstrip()
        s = raw.strip()

        if s.startswith("```"):
            in_code = not in_code
            i += 1
            continue
        if in_code:
            add_para(doc, raw, mono=True, indent=False)
            i += 1
            continue

        # 表格
        if s.startswith("|"):
            table_buf.append(split_row(s))
            i += 1
            if i >= len(lines) or not lines[i].strip().startswith("|"):
                rows = [r for r in table_buf if not all(re.fullmatch(r"-{2,}", c or "-") for c in r)]
                if rows:
                    add_table(doc, rows)
                    add_para(doc, "", size=6)
                table_buf = []
            continue

        # 行间公式：保留 LaTeX 原文
        if s.startswith("$$"):
            body = s
            if not (s.endswith("$$") and len(s) > 4):
                # 多行公式块
                buf = [s]
                i += 1
                while i < len(lines) and not lines[i].strip().endswith("$$"):
                    buf.append(lines[i].rstrip())
                    i += 1
                if i < len(lines):
                    buf.append(lines[i].rstrip())
                body = "\n".join(buf)
            add_para(doc, body, mono=True, indent=False)
            add_para(doc, "", size=6)
            i += 1
            continue

        if s.startswith("# "):
            add_para(doc, s[2:], size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        elif s.startswith("## "):
            add_para(doc, s[3:], size=14, bold=True, east="黑体")
        elif s.startswith("### "):
            add_para(doc, s[4:], size=12, bold=True, east="黑体")
        elif s.startswith(">"):
            add_para(doc, s.lstrip("> ").strip(), size=10.5, quote=True, indent=False)
        elif re.match(r"^([-*+]|\d+[.、)])\s+", s):
            add_para(doc, "• " + re.sub(r"^([-*+]|\d+[.、)])\s+", "", s), indent=False)
        elif s == "":
            pass
        else:
            add_para(doc, s, indent=True)
        i += 1

    doc.save(out_path)
    n_math = sum(1 for l in lines if l.strip().startswith("$$"))
    print(f"[OK] {src_path} -> {out_path}；行间公式 {n_math} 处（全部保留为 LaTeX 文本，无任何公式图片）")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
