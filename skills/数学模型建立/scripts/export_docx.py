# -*- coding: utf-8 -*-
"""把建模方案 Markdown 转成 DOCX（公式图片、表格、思路图、参考文献）。

用法：
    python export_docx.py --md <输入.md> [--out <输出.docx>] [--cache <公式缓存目录>]

前置要求：先运行 check_docx_toolchain.py 检测本机可用的解释器与工具链；
若检测结果显示缺少 python-docx / Pillow / matplotlib 或在线 LaTeX 兜底不可达，
按 references/文档导出与工具检测.md 的规则处理，必要时先询问用户，不要强行导出。
"""
from __future__ import annotations

import argparse
import hashlib
import re
import sys
import tempfile
import time
import urllib.parse
import urllib.request
from pathlib import Path

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

import matplotlib
matplotlib.use("Agg")
from matplotlib import mathtext  # noqa: E402
from PIL import Image  # noqa: E402
from docx import Document  # noqa: E402
from docx.enum.table import WD_TABLE_ALIGNMENT  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Cm, Pt, RGBColor  # noqa: E402

BS = chr(92)
CJK = re.compile("[\u3000-\u9fff\uff00-\uffef]")
TAG = re.compile(BS + BS + r"tag\{(\d+)\}")
DISPLAY = re.compile(r"\$\$(.*?)\$\$", re.DOTALL)


class MathRenderer:
    """优先用 matplotlib mathtext 本地渲染；失败则退回在线 LaTeX 服务（支持 cases 等复杂环境）。"""

    def __init__(self, cache: Path, dpi: int = 300):
        self.cache = cache
        self.cache.mkdir(parents=True, exist_ok=True)
        self.dpi = dpi
        self.stats = {"mathtext": 0, "online": 0, "fail": 0}

    def _key(self, latex: str, inline: bool) -> Path:
        h = hashlib.sha1((("i" if inline else "d") + latex).encode("utf-8")).hexdigest()[:16]
        return self.cache / f"{h}.png"

    def _mathtext(self, latex: str, target: Path) -> bool:
        try:
            mathtext.math_to_image("$" + latex + "$", str(target), dpi=self.dpi, format="png")
            return target.exists() and target.stat().st_size > 200
        except Exception:
            return False

    def _online(self, latex: str, target: Path, inline: bool) -> bool:
        prefix = BS + "dpi{300}" + (" " + BS + "inline " if inline else " ")
        url = "https://latex.codecogs.com/png.latex?" + urllib.parse.quote(prefix + latex)
        for attempt in range(2):
            try:
                req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
                with urllib.request.urlopen(req, timeout=40) as resp:
                    data = resp.read()
                if data[:8] == bytes([137, 80, 78, 71, 13, 10, 26, 10]):
                    target.write_bytes(data)
                    return True
            except Exception:
                time.sleep(1.5)
        return False

    def render(self, latex: str, inline: bool) -> Path | None:
        latex = normalize_math(latex)
        if not latex:
            return None
        target = self._key(latex, inline)
        if target.exists() and target.stat().st_size > 200:
            return target
        if self._mathtext(latex, target):
            self.stats["mathtext"] += 1
            return target
        if self._online(latex, target, inline):
            self.stats["online"] += 1
            return target
        self.stats["fail"] += 1
        print("     [WARN] 公式渲染失败:", latex[:60])
        return None


SKIP_LEAD = ("<!--", "---", "> 赛题", "> 交付", "> 引用", "> 配套", "> 结论")


def style_setup(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.35
    for name, size in (("Heading 1", 15), ("Heading 2", 13), ("Heading 3", 11.5)):
        st = doc.styles[name]
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.font.name = "微软雅黑"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        st.paragraph_format.space_before = Pt(12 if name == "Heading 1" else 9)
        st.paragraph_format.space_after = Pt(6)
    title = doc.styles["Title"]
    title.font.size = Pt(20)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.font.name = "微软雅黑"
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    for section in doc.sections:
        section.top_margin = Cm(2.4)
        section.bottom_margin = Cm(2.4)
        section.left_margin = Cm(2.6)
        section.right_margin = Cm(2.6)


def shade(cell, color: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def cell_borders(cell, color: str = "D9D9D9", size: int = 6) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    tcPr.append(borders)


def inline_runs(paragraph, text: str, renderer: MathRenderer, base_size: float = 10.5) -> None:
    """处理 **加粗**、`代码`、$行内公式$、[^n] 上标引用。"""
    text = re.sub(r"!\[[^\]]*\]\([^)]*\)", "", text)
    token = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\$[^$]+\$|\[\^\d+\])")
    for part in token.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = paragraph.add_run(part[2:-2]); r.bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = paragraph.add_run(part[1:-1]); r.font.name = "Consolas"; r.font.size = Pt(base_size - 0.5)
        elif part.startswith("$") and part.endswith("$") and len(part) > 2:
            png = renderer.render(part[1:-1], inline=True)
            if png:
                _add_picture_run(paragraph, png, target_pt=base_size + 1.5, max_width_in=5.9)
            else:
                paragraph.add_run(part[1:-1])
        elif re.fullmatch(r"\[\^\d+\]", part):
            r = paragraph.add_run(part[2:-1]); r.font.superscript = True
        else:
            paragraph.add_run(part)


def _add_picture_run(paragraph, png: Path, target_pt: float, max_width_in: float) -> None:
    with Image.open(png) as im:
        w_px, h_px = im.size
    dpi = 300.0
    width_in = w_px / dpi * (target_pt / 10.0)
    height_in = h_px / dpi * (target_pt / 10.0)
    if width_in > max_width_in:
        scale = max_width_in / width_in
        width_in *= scale
        height_in *= scale
    paragraph.add_run().add_picture(str(png), width=Cm(width_in * 2.54), height=Cm(height_in * 2.54))


def add_equation(doc: Document, latex: str, number: str | None, renderer: MathRenderer) -> None:
    png = renderer.render(latex, inline=False)
    table = doc.add_table(rows=1, cols=2)
    table.style = None  # 公式行不加边框，只做 居中公式 + 右侧编号 的排版
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Cm(14.0)
    table.columns[1].width = Cm(2.0)
    left, right = table.rows[0].cells
    left.width = Cm(14.0)
    right.width = Cm(2.0)
    para = left.paragraphs[0]
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if png:
        _add_picture_run(para, png, target_pt=13.0, max_width_in=5.6)
    else:
        para.add_run(latex)
    rp = right.paragraphs[0]
    rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if number:
        r = rp.add_run(f"({number})")
        r.font.size = Pt(10.5)
    doc.add_paragraph()


def split_cases(latex: str) -> tuple[str, list[str]]:
    """把 \\begin{cases} 拆成 (目标行, 各行约束文本)；中文标签转成普通文字，避免公式里出现中日韩字符。"""
    m = re.search(re.escape(BS + "begin{cases}") + r"(.*?)" + re.escape(BS + "end{cases}"), latex, re.DOTALL)
    if not m:
        return latex, []
    head = latex[: m.start()].replace(BS + "quad" + BS + "text{s.t.}", "").strip()
    head = re.sub(r"s\.t\.\s*$", "", head).strip()
    rows = []
    for raw in m.group(1).split(BS + BS):
        row = raw.strip().strip("&").strip()
        row = re.sub("^[" + BS + "[[][^]]*[]]]", "", row).strip()
        row = re.sub("[" + BS + "[[][0-9.]+[a-z]{2}[]]]", " ", row)
        row = re.sub(re.escape(BS + "text{") + r"([^}]*)\}", lambda mm: mm.group(1), row)
        row = re.sub(re.escape(BS + "mathrm{") + r"([^}]*)\}", lambda mm: mm.group(1), row)
        row = (row.replace(BS + "le", " <= ").replace(BS + "ge", " >= ")
                  .replace(BS + "in", " in ").replace(BS + "quad", " ")
                  .replace(BS + "qquad", " ").replace(BS + " ", " ").replace(BS + ",", " , "))
        row = re.sub(r"\s+", " ", row).strip(" ,")
        if row:
            rows.append(row)
    return head, rows


def normalize_math(latex):
    """把 mathtext 不支持的写法归一化，减少在线兜底调用。"""
    s = " ".join(latex.split())
    for token in (BS + "bigl", BS + "bigr", BS + "Bigl", BS + "Bigr", BS + "bigm", BS + "Bigm", BS + "!"):
        s = s.replace(token, "")
    s = s.replace(BS + "qquad", BS + "quad")
    s = s.replace(BS + "mathbb{1}", BS + "mathbf{1}")
    s = s.replace(BS + "dfrac", BS + "frac")
    s = s.replace(BS + "displaystyle", "")
    s = s.replace(BS + "tfrac", BS + "frac")
    for macro in ("sup", "inf", "arg", "det", "lim"):
        s = s.replace(BS + macro, BS + "mathrm{" + macro + "}")
    s = s.replace(BS + "operatorname{", BS + "mathrm{")
    s = s.replace(BS + "colon", ":").replace(BS + "mid", "|")
    s = re.sub(re.escape(BS + "ge") + "(?![a-zA-Z])", lambda m: BS + "geq", s)
    s = re.sub(re.escape(BS + "le") + "(?![a-zA-Z])", lambda m: BS + "leq", s)
    s = s.replace(BS + "iff", BS + "Leftrightarrow")
    pattern = re.compile(re.escape(BS + "text{") + "([A-Za-z0-9 ./+(),:-]*?)}")
    s = pattern.sub(lambda m: BS + "mathrm{" + m.group(1) + "}", s)
    # 中文标注：公式里不能出现中日韩字符（mathtext 与在线服务都无法稳定渲染），先换成英文等价词，再兜底删除
    for zh, en in (("拒收", "reject"), ("接收", "accept"), ("报废", "scrap"), ("拆解", "disassemble")):
        s = s.replace(zh, en)
    cjk_class = "[" + chr(0x3000) + "-" + chr(0x9FFF) + chr(0xFF00) + "-" + chr(0xFFEF) + "]"
    s = re.sub(re.escape(BS + "text{") + "[^}]*?" + cjk_class + "[^}]*?}", " ", s)
    return s.strip()


def convert(md: Path, out: Path, assets: Path, cache: Path) -> dict:
    text = md.read_text(encoding="utf-8")
    text = DISPLAY.sub(lambda m: "$" + "$" + m.group(1) + "$" + "$", text)
    lines = text.splitlines()
    doc = Document()
    style_setup(doc)
    renderer = MathRenderer(cache)
    stats = {"h1": 0, "h2": 0, "tables": 0, "equations": 0, "images": 0, "bullets": 0, "paras": 0}
    i = 0
    first_heading_done = False
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()
        if stripped.startswith("<!--") or stripped == "---":
            i += 1
            continue
        if not stripped:
            i += 1
            continue
        if stripped.startswith("# ") and not first_heading_done:
            doc.add_paragraph(stripped[2:], style="Title")
            first_heading_done = True
            i += 1
            continue
        if stripped.startswith("## "):
            doc.add_paragraph(stripped[3:], style="Heading 1"); stats["h1"] += 1
            i += 1
            continue
        if stripped.startswith("### "):
            doc.add_paragraph(stripped[4:], style="Heading 2"); stats["h2"] += 1
            i += 1
            continue
        if stripped.startswith("#### "):
            doc.add_paragraph(stripped[5:], style="Heading 3")
            i += 1
            continue
        if stripped.startswith("> "):
            p = doc.add_paragraph()
            r = p.add_run(stripped[2:])
            r.italic = True
            r.font.color.rgb = RGBColor(0x40, 0x40, 0x40)
            i += 1
            continue
        if stripped.startswith("![") and "](" in stripped:
            m = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", stripped)
            if m:
                img = Path(m.group(2))
                if not img.is_absolute():
                    img = (md.parent / img).resolve()
                if img.exists():
                    p = doc.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    _add_picture_run(p, img, target_pt=10.0, max_width_in=6.0)
                    if m.group(1):
                        cap = doc.add_paragraph()
                        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        cr = cap.add_run(m.group(1))
                        cr.font.size = Pt(9)
                        cr.font.color.rgb = RGBColor(0x40, 0x40, 0x40)
                    stats["images"] += 1
            i += 1
            continue
        if stripped.startswith("|") and i + 1 < len(lines) and set(lines[i + 1].strip()) <= set("|-: "):
            header = [c.strip() for c in stripped.strip("|").split("|")]
            rows = []
            j = i + 2
            while j < len(lines) and lines[j].strip().startswith("|"):
                rows.append([c.strip() for c in lines[j].strip().strip("|").split("|")])
                j += 1
            table = doc.add_table(rows=1, cols=len(header))
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for k, name in enumerate(header):
                cell = table.rows[0].cells[k]
                cell.text = ""
                para = cell.paragraphs[0]
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run_obj = para.add_run(name.replace("**", ""))
                run_obj.bold = True
                run_obj.font.size = Pt(9.5)
                run_obj.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                shade(cell, "2F5496")
                cell_borders(cell)
            for ridx, row in enumerate(rows):
                cells = table.add_row().cells
                for k in range(len(header)):
                    value = row[k] if k < len(row) else ""
                    cell = cells[k]
                    cell.text = ""
                    para = cell.paragraphs[0]
                    para.alignment = WD_ALIGN_PARAGRAPH.LEFT if len(value) > 14 else WD_ALIGN_PARAGRAPH.CENTER
                    inline_runs(para, value, renderer, base_size=9.0)
                    cell_borders(cell)
                    if ridx % 2 == 1:
                        shade(cell, "F2F5FA")
            doc.add_paragraph()
            stats["tables"] += 1
            i = j
            continue
        if stripped.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            inline_runs(p, stripped[2:], renderer)
            stats["bullets"] += 1
            i += 1
            continue
        if stripped.startswith("$$"):
            if stripped.endswith("$$") and len(stripped) > 4:
                body = stripped[2:-2]
                i += 1
            else:
                buf = [stripped[2:]]
                i += 1
                while i < len(lines) and "$$" not in lines[i]:
                    buf.append(lines[i])
                    i += 1
                if i < len(lines):
                    buf.append(lines[i].split("$$")[0])
                    i += 1
                body = chr(10).join(buf)
            latex = body.strip()
            numbers = TAG.findall(latex)
            number = numbers[0] if numbers else None
            latex = TAG.sub("", latex).strip()
            head, cases = split_cases(latex)
            if cases and CJK.search(head + " " + " ".join(cases)):
                add_equation(doc, head, number, renderer)  # 中文约束改为项目符号列表，避免公式里出现中日韩字符
            else:
                cases = []
                add_equation(doc, latex, number, renderer)  # 纯数学 cases 交给在线 LaTeX 渲染整式
            stats["equations"] += 1
            for row in cases:
                p = doc.add_paragraph(style="List Bullet")
                r = p.add_run(row)
                r.font.size = Pt(9.5)
            continue
        p = doc.add_paragraph()
        inline_runs(p, stripped, renderer)
        stats["paras"] += 1
        i += 1
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    stats["math"] = renderer.stats
    return stats


def main() -> int:
    ap = argparse.ArgumentParser(description="Markdown 转 DOCX（含公式与表格）")
    ap.add_argument("--md", required=True)
    ap.add_argument("--out", default=None, help="缺省为与 Markdown 同名的 .docx")
    ap.add_argument("--assets", default=None)
    ap.add_argument("--cache", default=None)
    args = ap.parse_args()
    md = Path(args.md).resolve()
    out = Path(args.out).resolve() if args.out else md.with_suffix(".docx")
    assets = Path(args.assets).resolve() if args.assets else md.parent / "diagrams"
    cache = Path(args.cache).resolve() if args.cache else Path(tempfile.gettempdir()) / "modeling_mathcache"
    stats = convert(md, out, assets, cache)
    print(f"[OK] {md.name} -> {out.name}")
    print("     " + ", ".join(f"{k}={v}" for k, v in stats.items() if k != "math"))
    print("     公式渲染来源:", stats["math"])
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
