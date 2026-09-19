# -*- coding: utf-8 -*-
"""把 data/clean 下的清洗结果汇总成 data/汇总数据.xlsx（每个表一个 sheet + 来源与字段 sheet）。

用法：
    python build_xlsx.py --clean "<交付目录>/data/clean" --out "<交付目录>/data/汇总数据.xlsx"
"""
from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

import pandas as pd  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

INVALID_SHEET = re.compile(r"[\\/*?:\[\]]")
HEADER_FILL = PatternFill("solid", fgColor="DCE6F1")


def sheet_name_of(path: Path, used: set[str]) -> str:
    name = INVALID_SHEET.sub("_", path.stem.replace(".clean", ""))[:28] or "Sheet"
    candidate = name
    i = 1
    while candidate in used:
        i += 1
        candidate = f"{name[:26]}_{i}"
    used.add(candidate)
    return candidate


def main() -> int:
    ap = argparse.ArgumentParser(description="汇总清洗数据为 xlsx")
    ap.add_argument("--clean", required=True, help="data/clean 目录")
    ap.add_argument("--out", required=True, help="输出的 xlsx 路径")
    ap.add_argument("--sources", default=None, help="sources.csv 路径，默认 <clean>/../sources.csv")
    ap.add_argument("--max-rows", type=int, default=1_048_000, help="单 sheet 行数上限保护")
    args = ap.parse_args()

    clean_dir = Path(args.clean)
    if not clean_dir.is_dir():
        print(f"[FAIL] 清洗目录不存在：{clean_dir}")
        return 1
    files = [p for p in sorted(clean_dir.glob("*.clean.csv"))]
    if not files:
        print(f"[FAIL] {clean_dir} 下没有 *.clean.csv")
        return 1

    out_path = Path(args.out)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    sources_path = Path(args.sources) if args.sources else clean_dir.parent / "sources.csv"

    used: set[str] = set()
    field_rows: list[dict] = []
    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        for path in files:
            df = pd.read_csv(path, dtype=object, encoding="utf-8-sig")
            if len(df) > args.max_rows:
                print(f"[WARN] {path.name} 行数 {len(df)} 超过上限，仅写入前 {args.max_rows} 行")
                df = df.head(args.max_rows)
            name = sheet_name_of(path, used)
            df.to_excel(writer, sheet_name=name, index=False)
            sheet = writer.sheets[name]
            for col_idx, col in enumerate(df.columns, 1):
                cell = sheet.cell(row=1, column=col_idx)
                cell.font = Font(bold=True)
                cell.fill = HEADER_FILL
                cell.alignment = Alignment(vertical="center")
                width = max(10, min(40, int(df[col].astype(str).str.len().head(200).max() or 10) + 2))
                sheet.column_dimensions[get_column_letter(col_idx)].width = width
            sheet.freeze_panes = "A2"
            for col in df.columns:
                field_rows.append({"表": name, "源文件": path.name, "字段": str(col),
                                   "非空数": int(df[col].notna().sum()), "总行数": int(len(df))})

        source_rows: list[dict] = []
        if sources_path.exists():
            source_rows = pd.read_csv(sources_path, dtype=object, encoding="utf-8-sig").to_dict("records")
        field_df = pd.DataFrame(field_rows) if field_rows else pd.DataFrame([{"表": "", "源文件": "", "字段": "", "非空数": "", "总行数": ""}])
        field_df.to_excel(writer, sheet_name="来源与字段", index=False, startrow=len(source_rows) + 4)
        sheet = writer.sheets["来源与字段"]
        sheet["A1"] = "数据来源清单（来自 data/sources.csv）"
        sheet["A1"].font = Font(bold=True, size=12)
        if source_rows:
            src_df = pd.DataFrame(source_rows)
            for col_idx, col in enumerate(src_df.columns, 1):
                cell = sheet.cell(row=2, column=col_idx, value=col)
                cell.font = Font(bold=True)
                cell.fill = HEADER_FILL
            for r, row in enumerate(source_rows, start=3):
                for c, col in enumerate(src_df.columns, 1):
                    sheet.cell(row=r, column=c, value=row.get(col))
        else:
            sheet["A2"] = "未找到 sources.csv：请先运行 fetch_data.py 登记来源。"
        start = len(source_rows) + 4
        sheet.cell(row=start, column=1, value="各表字段").font = Font(bold=True, size=12)
        for col_idx in range(1, 6):
            sheet.cell(row=start + 1, column=col_idx).fill = HEADER_FILL
        sheet.column_dimensions["A"].width = 22
        sheet.column_dimensions["B"].width = 32
        sheet.column_dimensions["C"].width = 24
        sheet.column_dimensions["D"].width = 16
        sheet.column_dimensions["E"].width = 12

    print(f"[OK] 已生成 {out_path}")
    print(f"     sheet：{'、'.join(sorted(used))}、来源与字段")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
