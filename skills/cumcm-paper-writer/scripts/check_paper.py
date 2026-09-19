#!/usr/bin/env python3
"""国赛论文机械自检：结构、编号、引用、匿名、篇幅与文件规格。

用法:
    python check_paper.py 论文正文.md [--docx 论文.docx] [--pdf 论文.pdf] [--out 格式自检表.md]

只做机械核对，不替代人工通读。发现 ❌ 时返回 1，仅 ⚠️ 时返回 0。
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

try:  # Windows console defaults to GBK; emoji output would crash
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")
except Exception:  # pragma: no cover
    pass

REQUIRED_SECTIONS = [
    ("问题重述", r"问题重述"),
    ("问题分析", r"问题分析"),
    ("模型假设", r"模型假设|模型的基本假设"),
    ("符号说明", r"符号说明|符号表"),
    ("模型建立与求解", r"模型建立|模型的建立"),
    ("模型检验", r"模型检验|误差分析|灵敏度|敏感性分析"),
    ("模型评价与推广", r"模型评价|模型的评价|模型推广"),
    ("参考文献", r"参考文献|References"),
    ("附录", r"附录"),
]

HARD_IDENTITY = re.compile(r"(参赛队号|报名号|队号[:：]|学号[:：]|姓名[:：]|指导教师|参赛学校|所在学校)")
SOFT_IDENTITY = re.compile(r"(大学|学院|校区|中学)")
PLACEHOLDER = re.compile(r"(TODO|待补|待填|XXX|××××|【\s*】|\?\?\?)")

MAX_FILE_MB = 20.0


def read_text(p: Path) -> str:
    return p.read_text(encoding="utf-8", errors="ignore")


def chinese_len(s: str) -> int:
    return len(re.findall(r"[\u4e00-\u9fff]", s))



RESIDUE_PATTERNS = {
    'bold-stars': r'\*\*',
    'heading-hash': r'(?m)^\s*#+\s',
    'table-dash': r'\|\s*-{3,}',
    'math-dollar': r'\$\$',
    'tag': r'\\tag\{',
    'image-syntax': r'!\[',
    'md-link': r'\]\(http',
    'backtick': r'`',
}


def check_markup_residue(doc_text):
    rows = []
    bad = {}
    for name, pat in RESIDUE_PATTERNS.items():
        n = len(re.findall(pat, doc_text))
        if n:
            bad[name] = n
    if bad:
        rows.append(('排版残留', 'Markdown 标记', '❌ ' + str(bad)))
    else:
        rows.append(('排版残留', 'Markdown 标记', '✅ 无残留'))
    return rows


def check_sections(text: str) -> list[tuple[str, str, str]]:
    rows = []
    for name, pat in REQUIRED_SECTIONS:
        hit = re.search(pat, text, re.IGNORECASE)
        rows.append(("结构", name, "✅" if hit else "❌ 缺章节"))
    return rows


def check_abstract(text: str) -> list[tuple[str, str, str]]:
    rows = []
    m = re.search(r"摘\s*要(.*?)(关键词|关键字)", text, re.S)
    if not m:
        rows.append(("摘要", "摘要正文", "❌ 未找到“摘要…关键词”结构"))
        return rows
    n = chinese_len(m.group(1))
    rows.append(("摘要", f"摘要字数 {n}",
                 "✅" if 700 <= n <= 1000 else "⚠️ 建议 700–1000 字（语料中位 831）"))
    kw = text[m.end() - 3: m.end() + 200] if m.end() >= 3 else ""
    kw = re.split(r"[\n。]", text[m.start(2): m.start(2) + 200])[0]
    cnt = len([x for x in re.split(r"[，,；;、\s]+", kw.replace("关键词", "").replace("关键字", "")) if len(x) >= 2])
    rows.append(("摘要", f"关键词 {cnt} 个", "✅" if 3 <= cnt <= 5 else "⚠️ 建议 3–5 个"))
    return rows


def check_numbering(text: str) -> list[tuple[str, str, str]]:
    rows = []
    lines = text.splitlines()

    eq = []
    for ln in lines:
        eq += [int(g) for g in re.findall(r"\\tag\{(\d{1,3})\}", ln)]
        m = re.search(r"[（(]\s*(\d{1,3})\s*[)）]\s*$", ln)
        if m:
            eq.append(int(m.group(1)))
    if eq:
        uniq = sorted(set(eq))
        gaps = [i for i in range(1, max(uniq) + 1) if i not in uniq]
        dup = len(eq) != len(uniq)
        rows.append(("编号", f"公式编号 共 {len(eq)} 处，最大 {max(uniq)}",
                     "✅" if not gaps and not dup else f"⚠️ 缺号 {gaps}；重复={dup}"))

        # 公式要逐式解释：语料中 69.5% 用“其中，”、27.7% 用“式中，”逐符号说明
        expl = len(re.findall(r"(?:式\s*中|其\s*中)\s*[：:，,]", text))
        ratio = expl / max(1, len(uniq))
        rows.append(("编号", f"公式解释 {expl} 处 / 公式 {len(uniq)} 个",
                     "✅ 逐式解释充分" if ratio >= 0.5
                     else "⚠️ 建议每个公式补“其中，… 为 …”的符号解释"))
    else:
        rows.append(("编号", "公式编号", "⚠️ 未检出编号公式"))

    # 题注可能在正文行首，也可能写在图片语法里：![图1 说明](x.png)；先剥掉 Markdown 装饰
    decorated = [re.sub(r"^[\s>*#\-|]*", "", ln) for ln in lines]
    for label, pat in (("图", r"^(?:图|!\[图)\s*(\d{1,3})"), ("表", r"^表\s*(\d{1,3})")):
        nums = [int(m.group(1)) for ln in decorated for m in [re.match(pat, ln)] if m]
        if not nums:
            rows.append(("编号", f"{label}题注", "⚠️ 未检出题注行"))
            continue
        gaps = [i for i in range(1, max(nums) + 1) if i not in set(nums)]
        rows.append(("编号", f"{label}题注 共 {len(nums)} 个，最大 {max(nums)}",
                     "✅" if not gaps else f"⚠️ 缺号 {gaps}"))
    return rows


def check_references(text: str) -> list[tuple[str, str, str]]:
    lines = text.splitlines()
    entries = [int(m.group(1)) for ln in lines for m in [re.match(r"^\s*\[(\d{1,3})\]\s*\S", ln)] if m]
    if not entries:
        return [("引用", "参考文献条目", "❌ 未检出参考文献条目")]

    ref_start = next((i for i, ln in enumerate(lines) if re.match(r"^\s*#*\s*参考文献", ln)), len(lines))
    body = "\n".join(lines[:ref_start])
    cited: set[int] = set()
    for m in re.finditer(r"\[([\d,\-–\s]+)\]", body):
        for part in re.split(r"[,，]", m.group(1)):
            part = part.strip()
            if not part:
                continue
            if re.match(r"^\d+$", part):
                cited.add(int(part))
            else:
                rng = re.match(r"^(\d+)\s*[-–]\s*(\d+)$", part)
                if rng:
                    cited.update(range(int(rng.group(1)), int(rng.group(2)) + 1))

    missing = sorted(set(entries) - cited)
    dangling = sorted(cited - set(entries))
    status = "✅"
    if missing:
        status = f"❌ 有未被引用的条目 {missing}"
    elif dangling:
        status = f"❌ 引用了不存在的条目 {dangling}"
    return [("引用", f"条目 {len(entries)} 条，正文引用 {len(cited)} 条", status)]


def check_anonymity(text: str) -> list[tuple[str, str, str]]:
    rows = []
    hard = [(i + 1, ln.strip()[:60]) for i, ln in enumerate(text.splitlines()) if HARD_IDENTITY.search(ln)]
    rows.append(("匿名", f"硬性身份信息 {len(hard)} 处",
                 "✅" if not hard else f"❌ {hard[:3]}"))

    soft = [(i + 1, ln.strip()[:60]) for i, ln in enumerate(text.splitlines())
            if SOFT_IDENTITY.search(ln) and not re.match(r"^\s*\[?\d*\]?\s*$", ln)]
    rows.append(("匿名", f"疑似机构名 {len(soft)} 处",
                 "✅" if not soft else f"⚠️ 逐条确认（参考文献作者单位可保留）：{soft[:3]}"))

    ph = [(i + 1, ln.strip()[:60]) for i, ln in enumerate(text.splitlines()) if PLACEHOLDER.search(ln)]
    rows.append(("完稿", f"占位符 {len(ph)} 处", "✅" if not ph else f"❌ 残留占位符 {ph[:3]}"))
    return rows


def check_files(md: Path, docx: Path | None, pdf: Path | None) -> list[tuple[str, str, str]]:
    rows = []
    total = chinese_len(read_text(md))
    rows.append(("篇幅", f"正文中文字数约 {total}", "✅" if total > 3000 else "⚠️ 偏短，确认是否完整"))

    for label, p in (("Word", docx), ("PDF", pdf)):
        if p is None:
            continue
        if not p.exists():
            rows.append(("文件", f"{label} {p.name}", "❌ 文件不存在"))
            continue
        mb = p.stat().st_size / 1024 / 1024
        rows.append(("文件", f"{label} {p.name} {mb:.2f} MB", "✅" if mb <= MAX_FILE_MB else "❌ 超过 20MB"))

    if docx is not None and docx.exists():
        try:
            from docx import Document
            d = Document(str(docx))
            s = d.sections[0]
            size_ok = abs(s.page_width.cm - 21.0) < 0.3 and abs(s.page_height.cm - 29.7) < 0.3
            rows.append(("版式", "纸张 A4" if size_ok else
                         f"纸张 {s.page_width.cm:.1f}×{s.page_height.cm:.1f} cm",
                         "✅" if size_ok else "⚠️ 非 A4"))
            rows.append(("版式", f"页边距 左右 {s.left_margin.cm:.2f}/{s.right_margin.cm:.2f} cm，上下 {s.top_margin.cm:.2f}/{s.bottom_margin.cm:.2f} cm",
                         "✅" if abs(s.left_margin.cm - 2.7) < 0.2 else "⚠️ 与模板（左2.7/上下2.54）不一致"))
            doc_text = "\n".join(p.text for p in d.paragraphs)
            hard = HARD_IDENTITY.findall(doc_text)
            rows.append(("匿名", f"Word 正文身份信息 {len(hard)} 处", "✅" if not hard else f"❌ {set(hard)}"))
        except ImportError:
            rows.append(("版式", "python-docx 不可用，跳过 Word 版式检查", "⚠️"))
        except Exception as e:  # pragma: no cover
            rows.append(("版式", f"Word 检查失败: {e}", "⚠️"))

    if pdf is not None and pdf.exists():
        try:
            from pypdf import PdfReader
            reader = PdfReader(str(pdf))
            n = len(reader.pages)
            # 规范只限制正文页数（≤30），附录页数不限：按“附录 A”标题所在页切分
            appendix_page = None
            for idx, page in enumerate(reader.pages):
                text = page.extract_text() or ""
                if re.search(r"附录\s*A\b|附录A", text):
                    appendix_page = idx + 1
                    break
            if appendix_page:
                body_pages = appendix_page - 2          # 本稿首页为摘要页
                rows.append(("篇幅", f"PDF 共 {n} 页：正文约 {body_pages} 页、附录自第 {appendix_page} 页起（正文限 ≤30 页，附录不限）",
                             "✅" if body_pages <= 30 else "⚠️ 正文超过 30 页"))
            else:
                rows.append(("篇幅", f"PDF 共 {n} 页（未找到附录起始页，按 n-3 估计正文）",
                             "⚠️ 超过 30 页" if n - 3 > 30 else "✅"))
        except ImportError:
            rows.append(("篇幅", "pypdf 不可用，跳过 PDF 页数检查", "⚠️"))
        except Exception as e:  # pragma: no cover
            rows.append(("篇幅", f"PDF 检查失败: {e}", "⚠️"))
    return rows


def main() -> int:
    ap = argparse.ArgumentParser(description="国赛论文机械自检")
    ap.add_argument("md", help="论文正文 Markdown/文本")
    ap.add_argument("--docx")
    ap.add_argument("--pdf")
    ap.add_argument("--out", help="自检表输出路径")
    args = ap.parse_args()

    md = Path(args.md)
    if not md.exists():
        print(f"文件不存在: {md}", file=sys.stderr)
        return 2
    text = read_text(md)
    # 附录里是源码与清单，可能出现 [0]、"图 x"、代码下标等字串；编号与引用检查只看正文
    # 只认真正的附录标题（"## 附录"），避免正文里"见附录 2"之类的行把正文截断
    _m_app = re.search(r"(?m)^#{1,3}\s*附\s*录", text)
    body_text = text[: _m_app.start()] if _m_app else text
    docx = Path(args.docx) if args.docx else None
    pdf = Path(args.pdf) if args.pdf else None

    rows: list[tuple[str, str, str]] = []
    rows += check_sections(text)
    rows += check_abstract(text)
    rows += check_numbering(body_text)
    rows += check_references(body_text)
    rows += check_anonymity(text)
    rows += check_files(md, docx, pdf)
    if docx is not None and docx.exists():
        try:
            from docx import Document as _D
            _doc = _D(str(docx))
            _txt = '\n'.join(p.text for p in _doc.paragraphs) + '\n' + '\n'.join(
                c.text for tb in _doc.tables for row in tb.rows for c in row.cells)
            # 附录里是源码：`#` 注释、`**` 幂运算都是合法代码，不能当排版残留；残留检查只看正文
            _m_app2 = re.search(r"\n\s*附\s*录", _txt)
            rows += check_markup_residue(_txt[: _m_app2.start()] if _m_app2 else _txt)
        except Exception:
            rows.append(('排版残留', 'Markdown 标记', '⚠️ 无法读取 Word 正文'))


    bad = [r for r in rows if r[2].startswith("❌")]
    warn = [r for r in rows if r[2].startswith("⚠️")]

    lines = ["# 论文格式自检表", "",
             f"- 稿件：`{md}`",
             f"- 结果：❌ {len(bad)} 项，⚠️ {len(warn)} 项，共 {len(rows)} 项", "",
             "| 类别 | 检查项 | 结论 |", "|---|---|---|"]
    lines += [f"| {a} | {b} | {c} |" for a, b, c in rows]
    lines += ["", "> 本表由 `check_paper.py` 生成，只做机械核对；❌ 必须修复或说明理由，⚠️ 需人工确认。"]
    report = "\n".join(lines) + "\n"

    if args.out:
        Path(args.out).write_text(report, encoding="utf-8")
        print(f"已写出: {args.out}")
    print(report)
    return 1 if bad else 0


if __name__ == "__main__":
    raise SystemExit(main())
