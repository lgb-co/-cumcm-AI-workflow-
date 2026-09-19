"""极简 Markdown → DOCX 转换：支持标题、段落、列表、表格、代码块、图片。

只依赖 python-docx；用于把附录/结果文档的 Markdown 源转成可提交的 Word 成品。
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

CODE_FONT = "Consolas"
BODY_FONT = "宋体"
HEAD_FONT = "黑体"


def _set_cjk_font(run, name: str) -> None:
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    from docx.oxml.ns import qn

    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        from docx.oxml import OxmlElement

        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        rfonts.set(qn(attr), name)


def _add_code_block(doc: Document, code: str) -> None:
    for line in code.splitlines() or [""]:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        run = paragraph.add_run(line if line.strip() else " ")
        run.font.size = Pt(8.5)
        _set_cjk_font(run, CODE_FONT)


def _add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for r, row in enumerate(rows):
        for c, cell in enumerate(row):
            if c >= len(rows[0]):
                continue
            paragraph = table.cell(r, c).paragraphs[0]
            run = paragraph.add_run(cell)
            run.font.size = Pt(9)
            if r == 0:
                run.bold = True
            _set_cjk_font(run, BODY_FONT)
    doc.add_paragraph()


def _add_image(doc: Document, path: Path, alt: str, width_cm: float = 13.0) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.add_picture(str(path), width=Cm(width_cm))
    if alt:
        caption = doc.add_paragraph()
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap_run = caption.add_run(alt)
        cap_run.font.size = Pt(9)
        cap_run.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
        _set_cjk_font(cap_run, BODY_FONT)


IMAGE_RE = re.compile(r"!\[(?P<alt>[^\]]*)\]\((?P<src>[^)]+)\)")
TABLE_SEP_RE = re.compile(r"^\s*\|?[\s:|-]+\|[\s:|-]*$")


def convert(md_text: str, docx_path: Path, base_dir: Path) -> Path:
    doc = Document()
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    style = doc.styles["Normal"]
    style.font.size = Pt(10.5)
    style.font.name = BODY_FONT
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attribute in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        rfonts.set(qn(attribute), BODY_FONT)

    lines = md_text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("```"):
            i += 1
            buffer: list[str] = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                buffer.append(lines[i])
                i += 1
            i += 1
            _add_code_block(doc, "\n".join(buffer))
            continue

        if stripped.startswith("|") and i + 1 < len(lines) and TABLE_SEP_RE.match(lines[i + 1]):
            rows: list[list[str]] = []
            header = [c.strip() for c in stripped.strip("|").split("|")]
            rows.append(header)
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            _add_table(doc, rows)
            continue

        image_match = IMAGE_RE.fullmatch(stripped)
        if image_match:
            src = image_match.group("src").strip()
            if not src.startswith("http"):
                image_path = (base_dir / src).resolve() if not Path(src).is_absolute() else Path(src)
                if image_path.exists():
                    _add_image(doc, image_path, image_match.group("alt"))
                else:
                    doc.add_paragraph(f"[缺图] {src}")
            i += 1
            continue

        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            text = stripped[level:].strip()
            heading = doc.add_heading("", level=min(level, 4))
            run = heading.add_run(text)
            _set_cjk_font(run, HEAD_FONT)
            i += 1
            continue

        if stripped in {"---", "***"}:
            i += 1
            continue

        bullet = re.match(r"^\s*[-*]\s+(.*)$", line)
        if bullet:
            paragraph = doc.add_paragraph(style="List Bullet")
            paragraph.add_run(bullet.group(1))
            i += 1
            continue

        numbered = re.match(r"^\s*\d+[.)]\s+(.*)$", line)
        if numbered:
            paragraph = doc.add_paragraph(style="List Number")
            paragraph.add_run(numbered.group(1))
            i += 1
            continue

        if stripped:
            paragraph = doc.add_paragraph()
            for chunk in re.split(r"(\*\*[^*]+\*\*)", stripped):
                if chunk.startswith("**") and chunk.endswith("**") and len(chunk) > 4:
                    run = paragraph.add_run(chunk[2:-2])
                    run.bold = True
                else:
                    run = paragraph.add_run(chunk)
                _set_cjk_font(run, BODY_FONT)
        i += 1

    docx_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(docx_path))
    return docx_path
